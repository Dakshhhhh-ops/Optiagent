"""Read CSV / Excel into plain tables and profile their columns.

Numbers are parsed here, by code, exactly once. The LLM never copies values.
"""

from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from pathlib import PurePath
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel

Cell = str | float | None

SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".xlsx")

# Limits that keep a single upload from exhausting memory or CPU.
MAX_ROWS = 100_000
MAX_COLUMNS = 1_000
MAX_SHEETS = 20
MAX_XLSX_UNCOMPRESSED_BYTES = 100_000_000
MISSING_TOKENS = frozenset({"", "-", "na", "n/a", "nan", "null", "none"})

_CURRENCY = re.compile(r"^[\s$€£₹¥]+|[\s$€£₹¥]+$")
_THOUSANDS = re.compile(r"^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$")


class Table(BaseModel):
    name: str
    columns: list[str]
    rows: list[list[Cell]]

    def column_index(self, column: str) -> int:
        try:
            return self.columns.index(column)
        except ValueError:
            raise KeyError(f"table {self.name!r} has no column {column!r}") from None

    def column_values(self, column: str) -> list[Cell]:
        i = self.column_index(column)
        return [row[i] if i < len(row) else None for row in self.rows]


class ColumnProfile(BaseModel):
    name: str
    dtype: Literal["number", "text", "empty"]
    non_empty: int
    distinct: int
    min: float | None = None
    max: float | None = None


class TableProfile(BaseModel):
    name: str
    row_count: int
    columns: list[ColumnProfile]
    preview: list[list[Cell]]


def is_missing(value: Cell) -> bool:
    return value is None or (isinstance(value, str) and value.strip().lower() in MISSING_TOKENS)


def parse_number(value: Cell) -> float | None:
    """Parse a cell as a number. Returns None if it is missing or not numeric.

    Accepts currency symbols and comma thousands separators ("$1,200.50").
    """
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        number = float(value)
        return number if math.isfinite(number) else None
    if value is None or is_missing(value):
        return None
    text = _CURRENCY.sub("", value.strip())
    if _THOUSANDS.match(text):
        text = text.replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None  # reject NaN and infinity


def as_text(value: Cell) -> str | None:
    if is_missing(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_header(header: list[Cell]) -> list[str]:
    names: list[str] = []
    for i, h in enumerate(header):
        base = as_text(h) or f"column_{i + 1}"
        name, n = base, 2
        while name in names:
            name, n = f"{base}_{n}", n + 1
        names.append(name)
    return names


def _drop_empty(rows: list[list[Cell]]) -> list[list[Cell]]:
    return [r for r in rows if any(not is_missing(c) for c in r)]


def read_csv(data: bytes, name: str) -> Table:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    records: list[list[Cell]] = []
    for record in csv.reader(io.StringIO(text), dialect):
        if len(records) > MAX_ROWS:
            raise ValueError(f"{name}: more than {MAX_ROWS:,} rows")
        if len(record) > MAX_COLUMNS:
            raise ValueError(f"{name}: more than {MAX_COLUMNS:,} columns")
        records.append(list(record))
    raw = _drop_empty(records)
    if not raw:
        return Table(name=name, columns=[], rows=[])
    header, *body = raw
    width = len(header)
    rows: list[list[Cell]] = [(r + [None] * width)[:width] for r in body]
    return Table(name=name, columns=_clean_header(list(header)), rows=rows)


def _check_xlsx_archive(data: bytes, name: str) -> None:
    """Reject zip bombs: .xlsx files that are small on disk but expand enormously."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{name}: not a valid .xlsx file") from exc
    expanded = sum(info.file_size for info in archive.infolist())
    if expanded > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError(
            f"{name}: expands to more than {MAX_XLSX_UNCOMPRESSED_BYTES // 1_000_000} MB"
        )


def read_excel(data: bytes, name: str) -> list[Table]:
    _check_xlsx_archive(data, name)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if len(workbook.worksheets) > MAX_SHEETS:
        raise ValueError(f"{name}: more than {MAX_SHEETS} sheets")
    tables = []
    for sheet in workbook.worksheets:
        raw: list[list[Cell]] = []
        for row in sheet.iter_rows(values_only=True):
            if len(raw) > MAX_ROWS:
                raise ValueError(f"{name}: sheet {sheet.title!r} has more than {MAX_ROWS:,} rows")
            if len(row) > MAX_COLUMNS:
                raise ValueError(
                    f"{name}: sheet {sheet.title!r} has more than {MAX_COLUMNS:,} columns"
                )
            raw.append(
                [c if isinstance(c, str | float | type(None)) else _excel_cell(c) for c in row]
            )
        raw = _drop_empty(raw)
        if not raw:
            continue
        header, *body = raw
        width = len(header)
        table_name = name if len(workbook.worksheets) == 1 else f"{name}:{sheet.title}"
        tables.append(
            Table(
                name=table_name,
                columns=_clean_header(header),
                rows=[(r + [None] * width)[:width] for r in body],
            )
        )
    workbook.close()
    return tables


def _excel_cell(value: object) -> Cell:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return float(value)
    return str(value)


def read_file(filename: str, data: bytes) -> list[Table]:
    path = PurePath(filename)
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        return [read_csv(data, path.name)]
    if suffix == ".xlsx":
        return read_excel(data, path.name)
    raise ValueError(f"unsupported file type {suffix!r}; use {', '.join(SUPPORTED_EXTENSIONS)}")


def profile_table(table: Table, preview_rows: int = 50) -> TableProfile:
    columns = []
    for column in table.columns:
        values = [v for v in table.column_values(column) if not is_missing(v)]
        numbers = [n for n in (parse_number(v) for v in values) if n is not None]
        if not values:
            dtype: Literal["number", "text", "empty"] = "empty"
        elif len(numbers) == len(values):
            dtype = "number"
        else:
            dtype = "text"
        columns.append(
            ColumnProfile(
                name=column,
                dtype=dtype,
                non_empty=len(values),
                distinct=len({as_text(v) for v in values}),
                min=min(numbers) if dtype == "number" else None,
                max=max(numbers) if dtype == "number" else None,
            )
        )
    return TableProfile(
        name=table.name,
        row_count=len(table.rows),
        columns=columns,
        preview=table.rows[:preview_rows],
    )
