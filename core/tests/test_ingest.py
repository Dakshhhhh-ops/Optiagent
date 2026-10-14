import io

import pytest
from openpyxl import Workbook

from optiagent_core.examples import list_sample_datasets, sample_dataset_files
from optiagent_core.ingest import (
    AssignmentMapping,
    KeyValue,
    LongLayout,
    MatrixLayout,
    Table,
    TransportationMapping,
    load_problem,
    profile_table,
    read_file,
    suggest_mapping,
)
from optiagent_core.ingest.tables import parse_number
from optiagent_core.pipeline import solve_problem
from optiagent_core.schema import AssignmentProblem, TransportationProblem


def csv(name: str, text: str) -> Table:
    return read_file(name, text.strip().encode())[0]


def sample_tables(sample_id: str) -> list[Table]:
    return [t for name, data in sample_dataset_files(sample_id) for t in read_file(name, data)]


# --------------------------------------------------------------------------- parsing


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("12", 12.0),
        (" 3.5 ", 3.5),
        ("$1,200.50", 1200.5),
        ("₹40", 40.0),
        ("-7", -7.0),
        (5, 5.0),
        ("", None),
        ("n/a", None),
        ("abc", None),
        ("1,2", None),  # ambiguous: not a thousands pattern
        ("nan", None),
    ],
)
def test_parse_number(raw: str | float, expected: float | None) -> None:
    assert parse_number(raw) == expected


def test_csv_sniffs_semicolons_and_strips_bom() -> None:
    table = read_file("x.csv", "﻿a;b\n1;2\n".encode())[0]
    assert table.columns == ["a", "b"]
    assert table.rows == [["1", "2"]]


def test_csv_dedupes_and_fills_headers() -> None:
    table = csv("x.csv", "id,cost,cost,\nA,1,2,3")
    assert table.columns == ["id", "cost", "cost_2", "column_4"]


def test_csv_skips_blank_rows_and_pads_short_rows() -> None:
    table = csv("x.csv", "a,b,c\n1,2\n,,\n4,5,6")
    assert table.rows == [["1", "2", None], ["4", "5", "6"]]


def test_excel_reads_each_sheet() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "supply"
    ws.append(["plant", "units"])
    ws.append(["P1", 10])
    ws2 = wb.create_sheet("demand")
    ws2.append(["city", "units"])
    ws2.append(["C1", 7.5])
    buf = io.BytesIO()
    wb.save(buf)

    tables = read_file("data.xlsx", buf.getvalue())
    assert [t.name for t in tables] == ["data.xlsx:supply", "data.xlsx:demand"]
    assert tables[0].rows == [["P1", 10.0]]
    assert tables[1].rows == [["C1", 7.5]]


def test_unsupported_extension() -> None:
    with pytest.raises(ValueError, match="unsupported"):
        read_file("data.pdf", b"")


def test_profile() -> None:
    profile = profile_table(csv("x.csv", "id,cost,note\nA,$5,hi\nB,7,\nC,1,"))
    by_name = {c.name: c for c in profile.columns}
    assert by_name["id"].dtype == "text"
    assert by_name["cost"].dtype == "number"
    assert (by_name["cost"].min, by_name["cost"].max) == (1, 7)
    assert by_name["note"].non_empty == 1
    assert profile.row_count == 3


# --------------------------------------------------------------------------- loading


SUPPLY = csv("supply.csv", "plant,cap\nP1,20\nP2,30")
DEMAND = csv("demand.csv", "city,need\nC1,25\nC2,25")
SUPPLY_KV = KeyValue(table="supply.csv", id_column="plant", value_column="cap")
DEMAND_KV = KeyValue(table="demand.csv", id_column="city", value_column="need")


def test_load_long_costs() -> None:
    costs = csv("costs.csv", "from,to,cost,limit\nP1,C1,2,\nP1,C2,4,10\nP2,C1,3,\nP2,C2,1,")
    mapping = TransportationMapping(
        costs=LongLayout(
            table="costs.csv",
            from_column="from",
            to_column="to",
            value_column="cost",
            capacity_column="limit",
        ),
        supply=SUPPLY_KV,
        demand=DEMAND_KV,
    )
    result = load_problem([costs, SUPPLY, DEMAND], mapping)
    assert isinstance(result.problem, TransportationProblem)
    routes = {(r.source, r.destination): (r.cost, r.capacity) for r in result.problem.routes}
    assert routes[("P1", "C2")] == (4, 10)
    assert routes[("P1", "C1")] == (2, None)
    assert solve_problem(result.problem).result.objective == pytest.approx(80)


def test_load_matrix_with_blocked_cells() -> None:
    costs = csv("costs.csv", "plant,C1,C2,total\nP1,2,x,99\nP2,3,1,99")
    mapping = TransportationMapping(
        costs=MatrixLayout(table="costs.csv", row_id_column="plant", ignore_columns=["total"]),
        supply=SUPPLY_KV,
        demand=DEMAND_KV,
    )
    result = load_problem([costs, SUPPLY, DEMAND], mapping)
    assert result.problem is not None
    assert len(result.problem.routes) == 3
    assert any("blocked" in i.message for i in result.issues if i.severity == "info")


def test_bad_number_is_reported_with_location() -> None:
    costs = csv("costs.csv", "plant,C1,C2\nP1,2,4\nP2,three,1")
    mapping = TransportationMapping(
        costs=MatrixLayout(table="costs.csv", row_id_column="plant"),
        supply=SUPPLY_KV,
        demand=DEMAND_KV,
    )
    result = load_problem([costs, SUPPLY, DEMAND], mapping)
    assert result.problem is None
    [error] = [i for i in result.issues if i.severity == "error"]
    assert (error.table, error.row, error.column) == ("costs.csv", 3, "C1")
    assert "three" in error.message


def test_unknown_source_in_costs_is_an_error() -> None:
    costs = csv("costs.csv", "plant,C1,C2\nP1,2,4\nP9,3,1")
    mapping = TransportationMapping(
        costs=MatrixLayout(table="costs.csv", row_id_column="plant"),
        supply=SUPPLY_KV,
        demand=DEMAND_KV,
    )
    result = load_problem([costs, SUPPLY, DEMAND], mapping)
    assert result.problem is None
    messages = [i.message for i in result.issues]
    assert any("'P9' has costs but no row" in m for m in messages)
    assert any("'P2' has no allowed pairs" in m for m in messages)


def test_duplicate_ids_are_errors() -> None:
    supply = csv("supply.csv", "plant,cap\nP1,20\nP1,30")
    costs = csv("costs.csv", "plant,C1\nP1,2")
    demand = csv("demand.csv", "city,need\nC1,5")
    mapping = TransportationMapping(
        costs=MatrixLayout(table="costs.csv", row_id_column="plant"),
        supply=SUPPLY_KV,
        demand=KeyValue(table="demand.csv", id_column="city", value_column="need"),
    )
    result = load_problem([costs, supply, demand], mapping)
    assert result.problem is None
    assert any("appears twice (rows 2 and 3)" in i.message for i in result.issues)


def test_missing_table_or_column_is_reported() -> None:
    mapping = TransportationMapping(
        costs=MatrixLayout(table="nope.csv", row_id_column="plant"),
        supply=SUPPLY_KV,
        demand=DEMAND_KV,
    )
    result = load_problem([SUPPLY, DEMAND], mapping)
    assert result.problem is None
    assert result.issues[0].message == "no table named 'nope.csv'"


def test_assignment_defaults_and_whole_numbers() -> None:
    prefs = csv("p.csv", "ta,lab,score\nA,L1,5\nB,L1,3")
    result = load_problem(
        [prefs],
        AssignmentMapping(
            sense="maximize",
            values=LongLayout(
                table="p.csv", from_column="ta", to_column="lab", value_column="score"
            ),
        ),
    )
    assert isinstance(result.problem, AssignmentProblem)
    assert [a.capacity for a in result.problem.agents] == [1, 1]
    assert [t.required for t in result.problem.tasks] == [1]

    labs = csv("labs.csv", "lab,needed\nL1,1.5")
    bad = load_problem(
        [prefs, labs],
        AssignmentMapping(
            values=LongLayout(
                table="p.csv", from_column="ta", to_column="lab", value_column="score"
            ),
            task_required=KeyValue(table="labs.csv", id_column="lab", value_column="needed"),
        ),
    )
    assert bad.problem is None
    assert any("whole number" in i.message for i in bad.issues)


# --------------------------------------------------------------------------- suggestions + samples


@pytest.mark.parametrize("sample", list_sample_datasets(), ids=lambda s: s.id)
def test_samples_suggest_load_and_solve(sample) -> None:  # type: ignore[no-untyped-def]
    tables = sample_tables(sample.id)
    suggestion = suggest_mapping(tables, sample.kind)
    assert suggestion.mapping is not None, suggestion.notes

    loaded = load_problem(tables, suggestion.mapping)
    assert loaded.problem is not None, loaded.issues

    report = solve_problem(loaded.problem)
    assert report.verified
    if sample.expected_objective is not None:
        assert report.result.objective == pytest.approx(sample.expected_objective)


def test_sunray_suggestion_shape() -> None:
    mapping = suggest_mapping(sample_tables("sunray"), "transportation").mapping
    assert isinstance(mapping, TransportationMapping)
    assert isinstance(mapping.costs, MatrixLayout)
    assert mapping.supply.table == "silos.csv"
    assert mapping.demand.table == "mills.csv"


def test_ta_labs_suggestion_maximizes_preferences() -> None:
    mapping = suggest_mapping(sample_tables("ta_labs"), "assignment").mapping
    assert isinstance(mapping, AssignmentMapping)
    assert mapping.sense == "maximize"
    assert mapping.task_required is not None
    assert mapping.task_required.table == "labs.csv"
    assert mapping.agent_capacity is None


def test_suggestion_uses_id_overlap_not_file_names() -> None:
    renamed = [
        t.model_copy(update={"name": n})
        for t, n in zip(sample_tables("sunray"), ["t1.csv", "t2.csv", "t3.csv"], strict=True)
    ]
    mapping = suggest_mapping(renamed, "transportation").mapping
    assert isinstance(mapping, TransportationMapping)
    # costs.csv -> t1, mills.csv -> t2, silos.csv -> t3 (sorted file order)
    assert (mapping.costs.table, mapping.demand.table, mapping.supply.table) == (
        "t1.csv",
        "t2.csv",
        "t3.csv",
    )
